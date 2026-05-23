import os
import tempfile
import unittest

import openpyxl

from core.report_export import (
    _BAD_FILL,
    _GOOD_FILL,
    _WARN_FILL,
    _WHITE,
    _mean_sd,
    _removal_fill,
    _safe_pct,
    build_experiment_report,
)


# ── Fixtures ───────────────────────────────────────────────────────────────────

def _exp():
    return {
        "name": "Test Experiment",
        "start_date": "2026-01-10",
        "end_date": "2026-03-20",
        "cec_spiked": "PFOS",
        "spike_conc": 100.0,
        "description": "Unit test run",
    }


def _daily_reading():
    return {
        "reading_date": "2026-02-01",
        "zone": "aeration",
        "do_mgl": 2.0,
        "ph": 7.2,
        "temperature_c": 22.0,
        "cod_mgl": None,
        "nh4_mgl": None,
        "no3_mgl": None,
        "so4_mgl": None,
        "po4_mgl": None,
        "mlss_mgl": 3000.0,
        "vss_mgl": None,
        "tn_mgl": None,
        "tp_mgl": None,
        "orp_mv": None,
        "turbidity_ntu": None,
        "our_mgl_h": None,
        "notes": "",
    }


def _cec_group(sp1=100.0, sp2=80.0, sp3=20.0, sp4=10.0):
    return {
        "date": "2026-02-01",
        "cec_name": "PFOS",
        "cec_type": "long-chain",
        "method": "LC-MS/MS",
        "sp1": sp1,
        "sp2": sp2,
        "sp3": sp3,
        "sp4": sp4,
    }


def _mb_result(pct=75.0):
    return {
        "calculation_date": "2026-02-01",
        "cec_name": "PFOS",
        "mass_in_mgday": 200.0,
        "mass_eff_mgday": 50.0,
        "mass_sludge_mgday": 100.0,
        "mass_transformed": 50.0,
        "removal_pct": pct,
        "q_eff_mlday": 290.0,
    }


def _sim_run():
    return {
        "run_date": "2026-02-01T10:00",
        "model_type": "Phase 1",
        "flow_rate_mlh": 320.0,
        "do_mgl": 2.0,
        "srt_days": 30.0,
        "ras_ratio": 1.0,
        "temp_c": 22.0,
        "ph": 7.5,
        "inf_cod": 300.0,
        "inf_nh4": 40.0,
        "inf_pfas": 50.0,
        "pfas_type": "short-chain",
        "pred_cod_removal": 80.0,
        "pred_nh4_removal": 90.0,
        "pred_tp_removal": 40.0,
        "pred_pfas_removal": 35.0,
    }


def _write_report(cec_groups=None, mb=None, sims=None, daily=None):
    fd, path = tempfile.mkstemp(suffix=".xlsx")
    os.close(fd)
    build_experiment_report(
        path,
        _exp(),
        daily or [_daily_reading()],
        cec_groups if cec_groups is not None else [_cec_group()],
        mb if mb is not None else [_mb_result()],
        sims if sims is not None else [_sim_run()],
    )
    return path


# ── Sheet structure ────────────────────────────────────────────────────────────

class ReportSheetStructureTests(unittest.TestCase):
    def test_creates_exactly_five_sheets(self):
        path = _write_report()
        try:
            wb = openpyxl.load_workbook(path)
            self.assertEqual(5, len(wb.sheetnames))
        finally:
            os.unlink(path)

    def test_sheet_names_match_expected(self):
        path = _write_report()
        try:
            wb = openpyxl.load_workbook(path)
            self.assertEqual(
                ["Summary", "Daily Readings", "CEC Readings", "Mass Balance", "Simulations"],
                wb.sheetnames,
            )
        finally:
            os.unlink(path)

    def test_empty_data_does_not_raise(self):
        path = _write_report(cec_groups=[], mb=[], sims=[], daily=[])
        try:
            wb = openpyxl.load_workbook(path)
            self.assertEqual(5, len(wb.sheetnames))
        finally:
            os.unlink(path)


# ── Summary sheet ──────────────────────────────────────────────────────────────

class SummarySheetTests(unittest.TestCase):
    def _all_values(self, ws, row_range=range(1, 20), col_range=range(1, 4)):
        return [str(ws.cell(r, c).value or "") for r in row_range for c in col_range]

    def test_experiment_name_appears_in_summary(self):
        path = _write_report()
        try:
            wb = openpyxl.load_workbook(path)
            values = self._all_values(wb["Summary"])
            self.assertTrue(any("Test Experiment" in v for v in values))
        finally:
            os.unlink(path)

    def test_cec_spiked_appears_in_summary(self):
        path = _write_report()
        try:
            wb = openpyxl.load_workbook(path)
            values = self._all_values(wb["Summary"])
            self.assertTrue(any("PFOS" in v for v in values))
        finally:
            os.unlink(path)

    def test_summary_key_statistics_section_present(self):
        path = _write_report()
        try:
            wb = openpyxl.load_workbook(path)
            values = self._all_values(wb["Summary"])
            self.assertTrue(any("Statistics" in v for v in values))
        finally:
            os.unlink(path)


# ── CEC sheet ──────────────────────────────────────────────────────────────────

class CECSheetTests(unittest.TestCase):
    def test_overall_removal_90_pct_formatted_correctly(self):
        # sp1=100, sp4=10 → (100-10)/100 × 100 = 90.0%
        path = _write_report(cec_groups=[_cec_group(sp1=100.0, sp4=10.0)])
        try:
            wb = openpyxl.load_workbook(path)
            ws = wb["CEC Readings"]
            overall = ws.cell(2, 12).value
            self.assertEqual("90.0%", overall)
        finally:
            os.unlink(path)

    def test_none_sp1_shows_dash_for_overall_removal(self):
        path = _write_report(cec_groups=[_cec_group(sp1=None, sp4=None)])
        try:
            wb = openpyxl.load_workbook(path)
            ws = wb["CEC Readings"]
            self.assertEqual("—", ws.cell(2, 12).value)
        finally:
            os.unlink(path)

    def test_multiple_cec_rows_written_in_order(self):
        groups = [
            _cec_group(sp1=100.0, sp4=10.0),
            _cec_group(sp1=200.0, sp4=20.0),
        ]
        path = _write_report(cec_groups=groups)
        try:
            wb = openpyxl.load_workbook(path)
            ws = wb["CEC Readings"]
            # Both rows should show 90.0% removal
            self.assertEqual("90.0%", ws.cell(2, 12).value)
            self.assertEqual("90.0%", ws.cell(3, 12).value)
        finally:
            os.unlink(path)

    def test_intermediate_stage_removals_computed(self):
        # sp1=100, sp2=80 → stage 1→2 removal = 20.0%
        path = _write_report(cec_groups=[_cec_group(sp1=100.0, sp2=80.0)])
        try:
            wb = openpyxl.load_workbook(path)
            ws = wb["CEC Readings"]
            stage12 = ws.cell(2, 9).value
            self.assertEqual("20.0%", stage12)
        finally:
            os.unlink(path)


# ── Mass Balance sheet ─────────────────────────────────────────────────────────

class MassBalanceSheetTests(unittest.TestCase):
    def test_removal_pct_cell_formatted_with_percent_sign(self):
        path = _write_report(mb=[_mb_result(pct=75.0)])
        try:
            wb = openpyxl.load_workbook(path)
            ws = wb["Mass Balance"]
            self.assertEqual("75.0%", ws.cell(2, 7).value)
        finally:
            os.unlink(path)

    def test_none_removal_pct_shows_dash(self):
        path = _write_report(mb=[_mb_result(pct=None)])
        try:
            wb = openpyxl.load_workbook(path)
            ws = wb["Mass Balance"]
            self.assertEqual("—", ws.cell(2, 7).value)
        finally:
            os.unlink(path)

    def test_mass_in_value_written_to_correct_column(self):
        path = _write_report(mb=[_mb_result()])
        try:
            wb = openpyxl.load_workbook(path)
            ws = wb["Mass Balance"]
            self.assertEqual(200.0, ws.cell(2, 3).value)
        finally:
            os.unlink(path)


# ── Simulations sheet ──────────────────────────────────────────────────────────

class SimulationsSheetTests(unittest.TestCase):
    def test_predicted_cod_removal_formatted_with_percent(self):
        path = _write_report(sims=[_sim_run()])
        try:
            wb = openpyxl.load_workbook(path)
            ws = wb["Simulations"]
            cod_cell = ws.cell(2, 13).value
            self.assertIn("%", str(cod_cell))
        finally:
            os.unlink(path)

    def test_none_prediction_shows_dash(self):
        run = _sim_run()
        run["pred_cod_removal"] = None
        path = _write_report(sims=[run])
        try:
            wb = openpyxl.load_workbook(path)
            ws = wb["Simulations"]
            self.assertEqual("—", ws.cell(2, 13).value)
        finally:
            os.unlink(path)


# ── _safe_pct helper ───────────────────────────────────────────────────────────

class SafePctTests(unittest.TestCase):
    def test_normal_removal_calculation(self):
        self.assertAlmostEqual(75.0, _safe_pct(100.0, 25.0), places=2)

    def test_none_influent_returns_none(self):
        self.assertIsNone(_safe_pct(None, 25.0))

    def test_none_effluent_returns_none(self):
        self.assertIsNone(_safe_pct(100.0, None))

    def test_zero_influent_returns_none(self):
        self.assertIsNone(_safe_pct(0.0, 10.0))

    def test_negative_influent_returns_none(self):
        self.assertIsNone(_safe_pct(-5.0, 10.0))

    def test_full_removal_gives_100_pct(self):
        self.assertAlmostEqual(100.0, _safe_pct(50.0, 0.0), places=2)

    def test_zero_removal_gives_zero_pct(self):
        self.assertAlmostEqual(0.0, _safe_pct(50.0, 50.0), places=2)


# ── _mean_sd helper ────────────────────────────────────────────────────────────

class MeanSdTests(unittest.TestCase):
    def test_multiple_values_returns_correct_mean(self):
        m, s = _mean_sd([80.0, 90.0, 70.0])
        self.assertAlmostEqual(80.0, m, places=2)

    def test_multiple_values_positive_stdev(self):
        _, s = _mean_sd([80.0, 90.0, 70.0])
        self.assertGreater(s, 0.0)

    def test_single_value_returns_zero_stdev(self):
        m, s = _mean_sd([75.0])
        self.assertEqual(75.0, m)
        self.assertEqual(0.0, s)

    def test_empty_list_returns_none_pair(self):
        m, s = _mean_sd([])
        self.assertIsNone(m)
        self.assertIsNone(s)

    def test_none_values_are_excluded_from_calculation(self):
        m, _ = _mean_sd([80.0, None, 90.0])
        self.assertAlmostEqual(85.0, m, places=2)

    def test_all_none_values_returns_none_pair(self):
        m, s = _mean_sd([None, None])
        self.assertIsNone(m)
        self.assertIsNone(s)


# ── _removal_fill helper ───────────────────────────────────────────────────────

class RemovalFillTests(unittest.TestCase):
    def test_high_removal_returns_green_fill(self):
        self.assertEqual(_GOOD_FILL, _removal_fill(75.0))

    def test_medium_removal_returns_amber_fill(self):
        self.assertEqual(_WARN_FILL, _removal_fill(45.0))

    def test_low_removal_returns_red_fill(self):
        self.assertEqual(_BAD_FILL, _removal_fill(10.0))

    def test_none_returns_white_fill(self):
        self.assertEqual(_WHITE, _removal_fill(None))

    def test_boundary_60_pct_is_green(self):
        self.assertEqual(_GOOD_FILL, _removal_fill(60.0))

    def test_boundary_30_pct_is_amber(self):
        self.assertEqual(_WARN_FILL, _removal_fill(30.0))

    def test_just_below_30_pct_is_red(self):
        self.assertEqual(_BAD_FILL, _removal_fill(29.9))


if __name__ == "__main__":
    unittest.main()
